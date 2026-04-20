# -*- coding: utf-8 -*-
"""Тесты записи Excel (с колонкой уверенности) и формирования лог-файла."""

from pathlib import Path
from unittest.mock import patch

import pytest
from openpyxl import load_workbook

from tn_parser import (
    FieldConfidence,
    ParsedRow,
    build_log_lines,
    write_excel,
    write_excel_safe,
    write_log,
    write_log_safe,
)
from tn_parser.excel import COLUMNS


def _sample_rows():
    high = ParsedRow(
        waybill="Транспортная накладная № 001",
        date="01.01.2024", number="001",
        shipper="ООО Ромашка", consignee="ООО Василёк",
        cargo="Мука", volume="500 шт", driver="Иванов И.И.",
        vehicle="KAMAZ\nА123ВС777", reception="г. Москва 01.01.2024",
        source="a.pdf", note="",
        confidence=FieldConfidence(
            date=1.0, number=0.9, shipper=0.9, consignee=0.9,
            cargo=0.9, volume=0.9, driver=0.9, vehicle=1.0, reception=0.9,
        ),
    )
    low = ParsedRow(
        waybill="Транспортная накладная",
        date="отсутствует", number="отсутствует",
        shipper="отсутствует", consignee="отсутствует",
        cargo="отсутствует", volume="отсутствует", driver="отсутствует",
        vehicle="отсутствует", reception="отсутствует",
        source="b.pdf", note="LOW_CONF",
        confidence=FieldConfidence(),
    )
    err = ParsedRow.empty_missing("c.pdf", note="ERROR: boom")
    return high, low, err


class TestExcel:
    def test_has_13_columns_including_confidence(self, tmp_path):
        rows = list(_sample_rows())
        out = tmp_path / "out.xlsx"
        write_excel(rows, str(out))

        wb = load_workbook(out)
        ws = wb.active
        headers = [ws.cell(row=1, column=i + 1).value for i in range(len(COLUMNS))]
        assert len(headers) == 13
        assert headers[-1] == "Уверенность, %"
        # Объём — между «Груз» и «Водитель».
        assert "Объём" in headers
        assert headers.index("Объём") == headers.index("Груз") + 1
        assert headers.index("Водитель") == headers.index("Объём") + 1

    def test_confidence_values(self, tmp_path):
        high, low, err = _sample_rows()
        out = tmp_path / "out.xlsx"
        write_excel([high, low, err], str(out))

        wb = load_workbook(out)
        ws = wb.active
        # Колонка confidence — 13-я.
        conf_col = len(COLUMNS)
        assert ws.cell(row=2, column=conf_col).value == 92  # high ≈ 0.92 → 92%
        assert ws.cell(row=3, column=conf_col).value == 0   # low
        assert ws.cell(row=4, column=conf_col).value == 0   # err

    def test_confidence_fill_colors(self, tmp_path):
        high, low, _ = _sample_rows()
        out = tmp_path / "out.xlsx"
        write_excel([high, low], str(out))
        wb = load_workbook(out)
        ws = wb.active
        conf_col = len(COLUMNS)
        high_fill = ws.cell(row=2, column=conf_col).fill.fgColor.rgb
        low_fill = ws.cell(row=3, column=conf_col).fill.fgColor.rgb
        assert "D9EAD3" in (high_fill or "").upper()
        assert "F4CCCC" in (low_fill or "").upper()

    def test_volume_cell_has_value(self, tmp_path):
        high, _, _ = _sample_rows()
        out = tmp_path / "out.xlsx"
        write_excel([high], str(out))
        wb = load_workbook(out)
        ws = wb.active
        # Объём — 7-я колонка.
        assert ws.cell(row=2, column=7).value == "500 шт"

    def test_vehicle_cell_wraps_on_newline(self, tmp_path):
        high, _, _ = _sample_rows()
        out = tmp_path / "out.xlsx"
        write_excel([high], str(out))
        wb = load_workbook(out)
        ws = wb.active
        # ТС — 9-я колонка, значение с переносом строки.
        cell = ws.cell(row=2, column=9)
        assert "\n" in (cell.value or "")
        assert cell.alignment.wrap_text  # перенос внутри ячейки включён


class TestLockedFileFallback:
    """Если основной файл заблокирован (открыт в Excel/блокноте), пишем
    под именем с меткой времени вместо того, чтобы падать."""

    def test_write_excel_safe_falls_back_on_permission_error(self, tmp_path):
        rows = list(_sample_rows())
        out = tmp_path / "locked.xlsx"

        # Первый вызов write_excel из write_excel_safe бросает PermissionError,
        # второй (на имя с меткой времени) проходит штатно.
        original = __import__("tn_parser.excel", fromlist=["write_excel"]).write_excel
        calls = []

        def fake_write_excel(rows, path):
            calls.append(path)
            if len(calls) == 1:
                raise PermissionError(13, "Permission denied", str(path))
            return original(rows, path)

        with patch("tn_parser.excel.write_excel", side_effect=fake_write_excel):
            actual = write_excel_safe(rows, str(out))

        # Возвращённый путь не равен исходному, но содержит базу и лежит в той же папке.
        assert actual != str(out)
        assert Path(actual).parent == tmp_path
        assert Path(actual).stem.startswith("locked_")
        assert Path(actual).exists()

    def test_write_excel_safe_ok_path_when_not_locked(self, tmp_path):
        rows = list(_sample_rows())
        out = tmp_path / "ok.xlsx"
        actual = write_excel_safe(rows, str(out))
        assert actual == str(out)
        assert out.exists()

    def test_write_log_safe_falls_back(self, tmp_path):
        out = tmp_path / "run.log"
        original = __import__("tn_parser.report", fromlist=["write_log"]).write_log
        calls = []

        def fake_write_log(path, lines):
            calls.append(path)
            if len(calls) == 1:
                raise PermissionError(13, "Permission denied", str(path))
            return original(path, lines)

        with patch("tn_parser.report.write_log", side_effect=fake_write_log):
            actual = write_log_safe(str(out), ["aaa", "bbb"])

        assert actual != str(out)
        assert Path(actual).exists()
        assert Path(actual).read_text(encoding="utf-8") == "aaa\nbbb\n"


class TestReport:
    def test_build_log_lines_structure(self):
        high, low, err = _sample_rows()
        lines = build_log_lines(
            input_path="/tmp/pdfs",
            output_path="/tmp/out.xlsx",
            elapsed_s=1.23,
            rows_by_file={"a.pdf": [high], "b.pdf": [low], "c.pdf": [err]},
        )
        text = "\n".join(lines)
        assert "Парсер транспортных накладных" in text
        assert "/tmp/pdfs" in text
        assert "[OK  ]" in text
        assert "[WARN]" in text
        assert "[ERR ]" in text
        assert "92%" in text
        # Для low confidence должны быть перечислены проблемные поля.
        assert "Проверить:" in text
        # Для ошибки должен быть виден её текст.
        assert "ERROR: boom" in text

    def test_write_log_creates_file(self, tmp_path):
        lines = ["строка один", "строка два"]
        target = tmp_path / "extraction.log"
        write_log(str(target), lines)
        content = target.read_text(encoding="utf-8")
        assert "строка один\nстрока два\n" == content
