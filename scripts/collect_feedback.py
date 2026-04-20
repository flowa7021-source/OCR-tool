"""Собирает правки оператора из Excel в корпус обучающих примеров.

Использование:

    python tools/collect_feedback.py out.xlsx
    python tools/collect_feedback.py out.xlsx --log feedback/corrections.jsonl

Как работает:

    1. Берём `out.xlsx` и рядом лежащий `out.xlsx.snapshot.json`
       (снапшот исходного вывода парсера — пишется автоматически при
        `write_excel_safe`).
    2. Читаем текущее содержимое Excel.
    3. Сравниваем построчно (по колонке «Источник файл») и выявляем
       поля, которые оператор изменил.
    4. Append'им правки в `feedback/corrections.jsonl`.

Результат: накопительный корпус правок, готовый для анализа слабых
мест парсера и построения few-shot промптов к LLM-fallback.
"""

from __future__ import annotations

import argparse
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from openpyxl import load_workbook  # noqa: E402

from src.tn_parser.feedback import (  # noqa: E402
    append_corrections,
    diff_rows,
    load_snapshot,
)

# Колонки Excel в порядке, заданном в tn_parser/excel.py::COLUMNS.
# Наполняем минимально нужный набор ключей для diff_rows.
_EXCEL_COLS = {
    1: "waybill",
    2: "date",
    3: "number",
    4: "shipper",
    5: "consignee",
    6: "cargo",
    7: "volume",
    8: "driver",
    9: "vehicle",
    10: "reception",
    11: "source",
    12: "note",
    13: "overall_confidence",
}


def read_excel_rows(xlsx_path: Path) -> list:
    wb = load_workbook(xlsx_path, data_only=True)
    ws = wb.active
    rows = []
    # Строка 1 — заголовки.
    for r in range(2, ws.max_row + 1):
        row_dict = {}
        for col, key in _EXCEL_COLS.items():
            val = ws.cell(row=r, column=col).value
            row_dict[key] = "" if val is None else str(val).strip()
        rows.append(row_dict)
    wb.close()
    return rows


def main(argv: list[str] | None = None) -> int:
    p = argparse.ArgumentParser(prog="ocr-cli parser collect-feedback")
    p.add_argument("xlsx", help="Путь к отредактированному Excel")
    p.add_argument(
        "--snapshot",
        help="Снапшот (по умолчанию <xlsx>.snapshot.json)",
        default=None,
    )
    p.add_argument(
        "--log",
        default="feedback/corrections.jsonl",
        help="Лог правок (по умолчанию feedback/corrections.jsonl)",
    )
    p.add_argument(
        "--dry-run",
        action="store_true",
        help="Показать правки, но не записывать в лог",
    )
    args = p.parse_args(argv)

    xlsx_path = Path(args.xlsx).resolve()
    if not xlsx_path.exists():
        print(f"ERROR: не найден {xlsx_path}", file=sys.stderr)
        return 1

    snapshot_path = Path(args.snapshot) if args.snapshot else Path(
        str(xlsx_path) + ".snapshot.json"
    )
    snapshot = load_snapshot(snapshot_path)
    if snapshot is None:
        print(
            f"ERROR: не найден снапшот {snapshot_path}. "
            "Снапшот пишется автоматически при сохранении Excel парсером. "
            "Если Excel был создан давно — восстановить не получится.",
            file=sys.stderr,
        )
        return 1

    snapshot_rows = snapshot.get("rows", [])
    corrected_rows = read_excel_rows(xlsx_path)

    corrections = diff_rows(snapshot_rows, corrected_rows)
    if not corrections:
        print("Правок не обнаружено.")
        return 0

    # Сводка по типам правок.
    by_field: dict = {}
    for c in corrections:
        by_field[c.field] = by_field.get(c.field, 0) + 1

    print(f"Обнаружено правок: {len(corrections)}")
    print("По полям:")
    for fld, n in sorted(by_field.items(), key=lambda x: -x[1]):
        print(f"  {fld:12s} {n}")

    if args.dry_run:
        print("\nDry-run: лог не обновлён.")
        for c in corrections[:10]:
            print(
                f"  [{c.kind}] {c.source} / {c.field}: "
                f"{c.original!r} → {c.corrected!r} "
                f"(conf={c.original_confidence:.2f})"
            )
        if len(corrections) > 10:
            print(f"  ... и ещё {len(corrections) - 10}")
        return 0

    log_path = Path(args.log).resolve()
    append_corrections(corrections, log_path)
    print(f"\nЗаписано в {log_path}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
