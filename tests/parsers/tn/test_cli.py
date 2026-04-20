"""Интеграционный тест CLI (`python -m src.tn_parser`).

Генерируем синтетический PDF через PyMuPDF из текстовой фикстуры и прогоняем
через CLI-функцию `main()`. Проверяем, что создаются Excel и лог.
"""

from pathlib import Path

import fitz
import pytest
from openpyxl import load_workbook

from src.tn_parser.__main__ import main as cli_main

FIXTURES = Path(__file__).parent / "fixtures"


def _make_pdf(text: str, dst: Path) -> None:
    """Создаёт PDF из текста: каждая логическая строка — на своей линии."""
    doc = fitz.open()
    for page_text in text.split("\f"):
        page = doc.new_page(width=595, height=842)  # A4
        y = 50
        for line in page_text.split("\n"):
            page.insert_text((40, y), line, fontsize=9, fontname="helv")
            y += 12
            if y > 800:
                page = doc.new_page(width=595, height=842)
                y = 50
    doc.save(str(dst))
    doc.close()


@pytest.fixture()
def pdf_folder(tmp_path: Path) -> Path:
    folder = tmp_path / "pdfs"
    folder.mkdir()
    std = (FIXTURES / "tn_standard.txt").read_text(encoding="utf-8")
    _make_pdf(std, folder / "a.pdf")
    _make_pdf(std, folder / "b.pdf")
    return folder


class TestCLI:
    def test_processes_folder_and_writes_outputs(self, pdf_folder: Path, tmp_path: Path):
        out = tmp_path / "result.xlsx"
        log = tmp_path / "result.log"
        rc = cli_main([
            str(pdf_folder),
            "--out", str(out),
            "--log", str(log),
            "--no-cache",
        ])
        assert rc == 0
        assert out.exists()
        assert log.exists()

        wb = load_workbook(out)
        ws = wb.active
        # Заголовок + 2 строки (по одной накладной на файл).
        assert ws.max_row >= 3
        # Последняя колонка — confidence (13-я после добавления «Объём»).
        assert ws.cell(row=1, column=13).value == "Уверенность, %"

        log_text = log.read_text(encoding="utf-8")
        assert "a.pdf" in log_text
        assert "b.pdf" in log_text

    def test_single_pdf_input(self, pdf_folder: Path, tmp_path: Path):
        single = pdf_folder / "a.pdf"
        out = tmp_path / "one.xlsx"
        rc = cli_main([str(single), "--out", str(out), "--no-cache"])
        assert rc == 0
        assert out.exists()
        # Лог должен лечь рядом с Excel.
        assert (tmp_path / "one.log").exists()

    def test_missing_input_returns_error(self, tmp_path: Path):
        rc = cli_main([str(tmp_path / "does_not_exist")])
        assert rc == 2

    def test_empty_folder_returns_error(self, tmp_path: Path):
        empty = tmp_path / "empty"
        empty.mkdir()
        rc = cli_main([str(empty)])
        assert rc == 1
