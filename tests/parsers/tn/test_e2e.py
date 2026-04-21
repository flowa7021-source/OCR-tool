"""E2E-тесты: повторяют пользовательский сценарий от PDF до Excel.

Пользователь на Windows:
    1. Кладёт PDF-файлы (машиночитаемые или OCR-слой) в папку.
    2. Запускает CLI `python -m src.tn_parser <папка>`.
    3. Получает `extraction.xlsx` и `extraction.log` рядом.
    4. Открывает Excel и видит по строке на каждую накладную.

Эти тесты гоняют ВЕСЬ конвейер: PDF → extract_best_text → normalize →
split_sections → extract_all → ParsedRow → Excel. Параметризуем на всех
golden-кейсах + OCR-фикстурах: если парсер ломается на реальном образце,
падает E2E, а не ручной фикс задним числом.

Переносимость: используем PyMuPDF для генерации синтетических PDF, что
позволяет тестам крутиться одинаково на Linux (CI) и Windows (dev/CI).
Пути — строго `pathlib.Path`, без зашитых разделителей.
"""

from __future__ import annotations

import json
from pathlib import Path

import fitz
import pytest
from openpyxl import load_workbook

from src.tn_parser import parse_text
from src.tn_parser.__main__ import main as cli_main
from src.tn_parser.excel import COLUMNS

FIXTURES = Path(__file__).parent / "fixtures"
GOLDEN = FIXTURES / "golden"


# --- Генерация синтетических PDF из текстовых фикстур ---------------------


def _make_pdf(text: str, dst: Path) -> None:
    """Текстовая фикстура → PDF с машиночитаемым текстовым слоем.

    Работает на любой ОС, не требует системных шрифтов — только PyMuPDF.
    Встроенный «helv» не содержит кириллицы (при рендере и последующем
    извлечении получаются точки). Поэтому используем встроенный CJK-шрифт
    `china-s` через `fitz.TextWriter` — он бандлится с PyMuPDF и содержит
    полный набор кириллических глифов, что даёт чистый round-trip.
    """
    # Ширина страницы намеренно сделана крупной (~2000pt): в реальных ТН
    # строки реквизитов бывают длиннее 200 символов, и при A4 + шрифте 9pt
    # хвост уходит за край листа — PyMuPDF при извлечении такой хвост
    # обрезает. Для E2E-прогонки нам нужно, чтобы ВЕСЬ текст гарантированно
    # попадал в экстракцию: «пакетный сканер» на проде обычно шлёт страницы
    # в «родной» геометрии документа, а не в A4.
    font = fitz.Font("china-s")
    doc = fitz.open()
    for page_text in text.split("\f"):
        page = doc.new_page(width=2000, height=2400)
        tw = fitz.TextWriter(page.rect)
        y = 50
        for line in page_text.split("\n"):
            if line:
                tw.append((40, y), line, font=font, fontsize=9)
            y += 12
            if y > 2350:
                tw.write_text(page)
                page = doc.new_page(width=2000, height=2400)
                tw = fitz.TextWriter(page.rect)
                y = 50
        tw.write_text(page)
    doc.save(str(dst))
    doc.close()


def _golden_cases() -> list[Path]:
    if not GOLDEN.exists():
        return []
    return sorted(GOLDEN.glob("*.expected.json"))


# --- E2E №1: один файл через CLI + проверка Excel и лога ------------------


class TestCliSingleFile:
    """Пользователь передал CLI одиночный PDF."""

    def test_produces_excel_with_correct_row(self, tmp_path: Path):
        raw = (FIXTURES / "tn_standard.txt").read_text(encoding="utf-8")
        pdf = tmp_path / "standard.pdf"
        _make_pdf(raw, pdf)

        out_xlsx = tmp_path / "out.xlsx"
        out_log = tmp_path / "out.log"

        rc = cli_main([
            str(pdf),
            "--out", str(out_xlsx),
            "--log", str(out_log),
            "--no-cache",
        ])
        assert rc == 0
        assert out_xlsx.exists(), "CLI must create an Excel file"
        assert out_log.exists(), "CLI must write a log next to Excel"

        wb = load_workbook(out_xlsx)
        ws = wb.active
        assert ws.cell(row=1, column=1).value == "Транспортная накладная"
        assert ws.cell(row=1, column=len(COLUMNS)).value == "Уверенность, %"

        # Значения в первой строке данных (row=2).
        waybill = ws.cell(row=2, column=1).value or ""
        number = ws.cell(row=2, column=3).value or ""
        assert "ТН-2024/00127" in waybill or number == "ТН-2024/00127"
        # conf% — int. Для standard ожидаем хотя бы 50% (новый контракт
        # строже — «Объём» в tn_standard.txt может быть не извлечён).
        conf = ws.cell(row=2, column=len(COLUMNS)).value
        assert isinstance(conf, int) and conf >= 50

        log_text = out_log.read_text(encoding="utf-8")
        assert "standard.pdf" in log_text


# --- E2E №2: папка с двумя PDF через CLI ----------------------------------


class TestCliFolderBatch:
    """Пользователь передал CLI папку — парсер должен обработать все PDF."""

    def test_folder_batch_creates_two_rows(self, tmp_path: Path):
        raw_a = (FIXTURES / "tn_standard.txt").read_text(encoding="utf-8")
        raw_b = (FIXTURES / "tn_ocr_tabular.txt").read_text(encoding="utf-8")
        folder = tmp_path / "inbox"
        folder.mkdir()
        _make_pdf(raw_a, folder / "01_standard.pdf")
        _make_pdf(raw_b, folder / "02_tabular.pdf")

        out_xlsx = tmp_path / "batch.xlsx"
        rc = cli_main([
            str(folder),
            "--out", str(out_xlsx),
            "--no-cache",
        ])
        assert rc == 0
        assert out_xlsx.exists()

        wb = load_workbook(out_xlsx)
        ws = wb.active
        # Заголовок + 2 строки данных.
        assert ws.max_row >= 3
        # «Источник файл» — ищем по заголовку (позиция сдвинулась
        # после добавления реквизитов в апреле 2026).
        from src.tn_parser.excel import COLUMNS
        source_idx = [c[0] for c in COLUMNS].index("Источник файл") + 1
        sources = {
            ws.cell(row=r, column=source_idx).value
            for r in range(2, ws.max_row + 1)
        }
        assert "01_standard.pdf" in sources
        assert "02_tabular.pdf" in sources

    def test_golden_case_volume_field_in_excel(self, tmp_path: Path):
        """Объём из golden-кейса tn_ocr_tabular (720 шт) доезжает до Excel."""
        raw = (FIXTURES / "tn_ocr_tabular.txt").read_text(encoding="utf-8")
        folder = tmp_path / "in"
        folder.mkdir()
        _make_pdf(raw, folder / "case.pdf")
        out_xlsx = tmp_path / "vol.xlsx"
        rc = cli_main([str(folder), "--out", str(out_xlsx), "--no-cache"])
        assert rc == 0
        wb = load_workbook(out_xlsx)
        ws = wb.active
        # «Объём» — ищем по заголовку (позиция 7 → 13 после добавления
        # 6 реквизитов отправителя/получателя в апреле 2026).
        from src.tn_parser.excel import COLUMNS
        volume_idx = [c[0] for c in COLUMNS].index("Объём") + 1
        volume_values = {
            ws.cell(row=r, column=volume_idx).value
            for r in range(2, ws.max_row + 1)
        }
        assert any(v and "720 шт" in v for v in volume_values), (
            f"expected '720 шт' in volume column, got {volume_values!r}"
        )


# --- E2E №3: все golden-кейсы через полный конвейер PDF → ParsedRow -------


@pytest.mark.parametrize(
    "expected_path",
    _golden_cases(),
    ids=[p.stem.replace(".expected", "") for p in _golden_cases()],
)
def test_golden_case_via_pdf_pipeline(expected_path: Path, tmp_path: Path):
    """Каждый golden-кейс синтезируется в PDF, прогоняется через весь
    конвейер (в т. ч. PyMuPDF-извлечение). Затем сверяется по contains-
    правилам из expected.json.

    Отличие от `test_golden.py`: там мы парсим ТЕКСТ фикстуры, тут —
    экспортируем её в PDF и читаем обратно. Это ловит регрессии в слое
    `layout.extract_best_text` и в нормализации, которые text-only тесты
    пропустили бы.
    """
    case = expected_path.stem.replace(".expected", "")
    txt_path = expected_path.with_name(f"{case}.txt")
    if not txt_path.exists():
        pytest.skip(f"raw text missing: {txt_path.name}")

    raw = txt_path.read_text(encoding="utf-8")
    pdf_path = tmp_path / f"{case}.pdf"
    _make_pdf(raw, pdf_path)

    from src.tn_parser.core import extract_raw_text  # локальный импорт для ясности
    text = extract_raw_text(str(pdf_path))
    rows = parse_text(text, f"{case}.pdf")

    spec = json.loads(expected_path.read_text(encoding="utf-8"))
    expected_rows = spec.get("rows") or []
    assert len(rows) == len(expected_rows), (
        f"[{case}] parsed {len(rows)} rows via PDF, expected {len(expected_rows)}"
    )

    for _i, (row, exp) in enumerate(zip(rows, expected_rows, strict=False)):
        for field in ("number", "date", "shipper", "consignee",
                      "cargo", "driver", "vehicle", "reception"):
            if field not in exp:
                continue
            value = getattr(row, field)
            rule = exp[field]
            if isinstance(rule, dict):
                for needle in rule.get("contains", []) or []:
                    assert needle in value, (
                        f"[{case} via PDF] {field}: "
                        f"expected to contain {needle!r}, got {value!r}"
                    )
                for needle in rule.get("not_contains", []) or []:
                    assert needle not in value, (
                        f"[{case} via PDF] {field}: "
                        f"must NOT contain {needle!r}, got {value!r}"
                    )


# --- E2E №4: локкнутый Excel — CLI пишет под соседним именем --------------


class TestCliExcelLocked:
    """Пользователь на Windows запустил парсинг, а целевой .xlsx открыт
    в Excel и заблокирован. CLI должен сохранить в файл с таймстампом и
    вернуть 0, а не упасть."""

    def test_fallback_when_target_is_locked(self, tmp_path: Path, monkeypatch):
        raw = (FIXTURES / "tn_standard.txt").read_text(encoding="utf-8")
        pdf = tmp_path / "standard.pdf"
        _make_pdf(raw, pdf)

        out_xlsx = tmp_path / "locked.xlsx"

        # Имитируем «файл занят другим процессом»: первая запись падает с
        # PermissionError, безопасный вариант в write_excel_safe должен
        # дописать под именем с таймстампом.
        from src.tn_parser import excel as excel_mod
        real_write_excel = excel_mod.write_excel
        calls = {"n": 0}

        def flaky_write(rows, path):
            calls["n"] += 1
            if calls["n"] == 1:
                raise PermissionError("simulated: file open in Excel")
            return real_write_excel(rows, path)

        monkeypatch.setattr(excel_mod, "write_excel", flaky_write)

        rc = cli_main([str(pdf), "--out", str(out_xlsx), "--no-cache"])
        assert rc == 0
        # Оригинал НЕ создан (первый write упал), но рядом должен быть файл
        # с таймстампом в имени.
        siblings = list(tmp_path.glob("locked_*.xlsx"))
        assert siblings, "expected a timestamped fallback .xlsx next to the locked one"
