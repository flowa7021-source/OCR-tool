"""Unit + integration tests for ``ExportManager.export_excel``.

The contract under test is the complete round-trip from a finished
OCR job with ``result.parsed`` populated to an on-disk ``.xlsx`` file
that the ТН / УПД CLI's Excel output can be diff'd against byte-for-
byte. The goal is an integration that:

* Refuses to export when there are no parsed rows (silently writing
  an empty workbook would hide a profile / extract-config bug).
* Fails fast with a typed :class:`ExportError` when parser
  dependencies (openpyxl, rapidfuzz) are missing in a trimmed
  install — not a raw ``ImportError`` that the UI can't translate
  into a user message.
* Writes the ``.xlsx`` + ``.log`` + ``.snapshot.json`` trio in one
  call, matching the standalone parser's CLI layout so downstream
  feedback tooling (``scripts/collect_feedback.py``) keeps working
  regardless of whether the file came from the CLI or the GUI.
* Populates ``job_result.parsed.snapshot_path`` so callers don't
  have to re-derive it.
"""

from __future__ import annotations

from pathlib import Path

import pytest

from src.application.export_manager import ExportError, ExportManager
from src.core.models import (
    JobResult,
    PageResult,
    ParsedDocument,
)
from src.shared.types import ExportFormat, JobStatus
from src.tn_parser import parse_text
from src.tn_parser.normalize import normalize_for_sections

_FIXTURES = Path(__file__).resolve().parents[1] / "parsers" / "tn" / "fixtures"
_TABULAR_FIXTURE = _FIXTURES / "tn_ocr_tabular.txt"


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------


def _parsed_from_fixture() -> ParsedDocument:
    """Build a ParsedDocument by running the real parser on a fixture.

    Using the real parser (not a hand-crafted stub) keeps this test
    honest: if the parser's ``to_json_dict`` shape drifts, the
    ``from_json_dict`` round-trip in ``export_excel`` will catch it
    here before it silently corrupts a user's Excel output.
    """
    raw = _TABULAR_FIXTURE.read_text(encoding="utf-8")
    rows = parse_text(normalize_for_sections(raw), "tn_ocr_tabular.pdf")
    assert rows, "fixture must produce at least one parsed row"
    return ParsedDocument(
        rows=[r.to_json_dict() for r in rows],
        overall_confidence=round(
            sum(r.confidence.overall() for r in rows) / len(rows), 3,
        ),
    )


def _job_result(parsed: ParsedDocument | None) -> JobResult:
    return JobResult(
        job_id="test-job",
        status=JobStatus.COMPLETED,
        input_path="/tmp/tn_ocr_tabular.pdf",
        output_path="/tmp/out.pdf",
        pages=[PageResult(page_number=1, text="stub")],
        total_time_sec=1.23,
        parsed=parsed,
    )


# ---------------------------------------------------------------------------
# Gate behaviour
# ---------------------------------------------------------------------------


def test_export_excel_raises_when_parsed_is_none(tmp_path: Path) -> None:
    """Profile without extract → parsed is None → explicit error."""
    mgr = ExportManager()
    result = _job_result(parsed=None)
    with pytest.raises(ExportError) as exc_info:
        mgr.export_excel(result, tmp_path / "out.xlsx")
    # Error message should hint at the profile flag, not just say
    # "null pointer" — the user is the target audience.
    assert "extract.enabled" in str(exc_info.value)


def test_export_excel_raises_when_rows_empty(tmp_path: Path) -> None:
    """Parser ran but produced no rows → still an error (not empty file)."""
    mgr = ExportManager()
    result = _job_result(parsed=ParsedDocument(rows=[], overall_confidence=0.0))
    with pytest.raises(ExportError):
        mgr.export_excel(result, tmp_path / "out.xlsx")


def test_dispatch_via_format_enum(tmp_path: Path) -> None:
    """``export(format=ExportFormat.EXCEL)`` routes to export_excel."""
    mgr = ExportManager()
    result = _job_result(parsed=_parsed_from_fixture())
    written = mgr.export(result, tmp_path / "dispatched.xlsx", ExportFormat.EXCEL)
    assert written.exists()
    assert written.suffix == ".xlsx"


# ---------------------------------------------------------------------------
# Happy path — integration with the real parser + openpyxl
# ---------------------------------------------------------------------------


def test_export_excel_writes_xlsx_log_and_snapshot(tmp_path: Path) -> None:
    """One call produces three files: xlsx, log, snapshot.

    Matches the CLI's on-disk layout — downstream tooling
    (``scripts/collect_feedback.py``) expects all three to sit next
    to each other with related names.
    """
    from openpyxl import load_workbook

    mgr = ExportManager()
    parsed = _parsed_from_fixture()
    result = _job_result(parsed=parsed)
    out = tmp_path / "extraction.xlsx"

    written = mgr.export_excel(result, out)

    # xlsx written at the requested path (no lock → no fallback name).
    assert written == out
    assert out.exists() and out.stat().st_size > 0

    # Log sidecar — same stem as the xlsx.
    log_path = out.with_suffix(".log")
    assert log_path.exists(), "log file must be written next to xlsx"
    log_text = log_path.read_text(encoding="utf-8")
    # Report header and at least one row line ([OK  ], [WARN], etc).
    assert "Парсер транспортных накладных" in log_text
    assert "tn_ocr_tabular.pdf" in log_text

    # Snapshot sidecar — used by collect_feedback.py.
    snapshot_path = Path(str(out) + ".snapshot.json")
    assert snapshot_path.exists(), "snapshot.json sidecar must be written"
    assert result.parsed.snapshot_path == str(snapshot_path)

    # Contents verifiable by openpyxl. Последняя колонка = «Уверенность, %»,
    # позиция определяется по ``COLUMNS`` (19 после апреля 2026, было 13).
    from src.tn_parser.excel import COLUMNS
    wb = load_workbook(out)
    ws = wb.active
    assert ws.cell(row=1, column=1).value == "Транспортная накладная"
    assert ws.cell(row=1, column=len(COLUMNS)).value == "Уверенность, %"
    # Row 2 is the one parsed row from tn_ocr_tabular.txt: number 7145/Б.
    row2 = [
        ws.cell(row=2, column=c).value or ""
        for c in range(1, len(COLUMNS) + 1)
    ]
    assert any("7145" in str(v) for v in row2), f"unexpected row: {row2!r}"


def test_export_excel_uses_fallback_name_when_locked(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    """When write_excel fails with PermissionError, write_excel_safe
    falls back to a timestamped name, and the returned path must
    reflect that — snapshot_path must follow the fallback too.
    """
    mgr = ExportManager()
    parsed = _parsed_from_fixture()
    result = _job_result(parsed=parsed)
    out = tmp_path / "locked.xlsx"

    import src.tn_parser.excel as excel_mod
    original_write = excel_mod.write_excel
    calls: list[str] = []

    def fake_write_excel(rows, path):  # noqa: ANN001
        calls.append(path)
        if len(calls) == 1:
            raise PermissionError(13, "Permission denied", str(path))
        return original_write(rows, path)

    monkeypatch.setattr(excel_mod, "write_excel", fake_write_excel)

    written = mgr.export_excel(result, out)
    assert written != out, "fallback path must differ from original"
    assert written.parent == tmp_path
    assert written.stem.startswith("locked_"), f"unexpected name: {written.name}"
    assert written.exists()
    # Snapshot path must match the actual xlsx, not the requested one —
    # else feedback tooling looks for snapshot next to a file that
    # doesn't exist.
    assert result.parsed.snapshot_path == str(written) + ".snapshot.json"
    assert Path(result.parsed.snapshot_path).exists()
