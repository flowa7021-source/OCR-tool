"""Tests for the ExportManager: TXT, DOCX, PDF passthrough, dispatcher."""

from __future__ import annotations

from pathlib import Path

import pytest

from src.application.export_manager import ExportManager
from src.core.models import JobResult, PageResult
from src.shared.types import ExportFormat, JobStatus


def _sample_result(output_path: Path) -> JobResult:
    return JobResult(
        job_id="j1",
        status=JobStatus.COMPLETED,
        input_path="/tmp/in.pdf",
        output_path=str(output_path),
        pages=[
            PageResult(page_number=1, text="Первая страница.\n\nВторой абзац.", mean_confidence=92.5),
            PageResult(page_number=2, text="Hello world.", mean_confidence=88.1),
            PageResult(page_number=3, error="Timeout"),
        ],
        total_time_sec=5.5,
    )


class TestExportTxt:
    def test_writes_file(self, tmp_path: Path) -> None:
        result = _sample_result(tmp_path / "out.pdf")
        target = tmp_path / "out.txt"
        written = ExportManager().export_txt(result, target)
        assert written == target
        assert target.exists()

    def test_contains_all_page_text(self, tmp_path: Path) -> None:
        result = _sample_result(tmp_path / "out.pdf")
        target = tmp_path / "out.txt"
        ExportManager().export_txt(result, target)
        body = target.read_text(encoding="utf-8")
        assert "Первая страница." in body
        assert "Второй абзац." in body
        assert "Hello world." in body
        # Page headers present
        assert "Page 1" in body
        assert "Page 2" in body
        assert "Page 3" in body

    def test_error_marker_preserved(self, tmp_path: Path) -> None:
        result = _sample_result(tmp_path / "out.pdf")
        target = tmp_path / "out.txt"
        ExportManager().export_txt(result, target)
        body = target.read_text(encoding="utf-8")
        assert "[ERROR: Timeout]" in body

    def test_unicode_roundtrip(self, tmp_path: Path) -> None:
        result = _sample_result(tmp_path / "out.pdf")
        target = tmp_path / "out.txt"
        ExportManager().export_txt(result, target)
        # Ensure we didn't accidentally double-encode Cyrillic
        raw = target.read_bytes()
        assert "Первая".encode() in raw

    def test_bom_when_utf8_sig(self, tmp_path: Path) -> None:
        result = _sample_result(tmp_path / "out.pdf")
        target = tmp_path / "out.txt"
        ExportManager().export_txt(result, target, encoding="utf-8-sig")
        raw = target.read_bytes()
        assert raw.startswith(b"\xef\xbb\xbf")

    def test_creates_parent_dir(self, tmp_path: Path) -> None:
        result = _sample_result(tmp_path / "out.pdf")
        target = tmp_path / "sub" / "deep" / "out.txt"
        ExportManager().export_txt(result, target)
        assert target.exists()


class TestExportDocx:
    def test_writes_file(self, tmp_path: Path) -> None:
        pytest.importorskip("docx")
        result = _sample_result(tmp_path / "out.pdf")
        target = tmp_path / "out.docx"
        written = ExportManager().export_docx(result, target)
        assert written == target
        assert target.exists()
        # DOCX is a ZIP archive
        assert target.read_bytes().startswith(b"PK")

    def test_contains_page_headings(self, tmp_path: Path) -> None:
        pytest.importorskip("docx")
        from docx import Document

        result = _sample_result(tmp_path / "out.pdf")
        target = tmp_path / "out.docx"
        ExportManager().export_docx(result, target)

        doc = Document(str(target))
        all_text = "\n".join(p.text for p in doc.paragraphs)
        assert "Page 1" in all_text
        assert "Первая страница." in all_text
        assert "Hello world." in all_text

    def test_error_page_marked(self, tmp_path: Path) -> None:
        pytest.importorskip("docx")
        from docx import Document

        result = _sample_result(tmp_path / "out.pdf")
        target = tmp_path / "out.docx"
        ExportManager().export_docx(result, target)
        doc = Document(str(target))
        all_text = "\n".join(p.text for p in doc.paragraphs)
        assert "[ERROR: Timeout]" in all_text


class TestExportPdf:
    def test_copy_to_new_location(self, tmp_path: Path) -> None:
        """Save As to a fresh path copies the searchable PDF."""
        from src.application.export_manager import ExportManager

        source = tmp_path / "searchable.pdf"
        source.write_bytes(b"%PDF-1.7\nbody bytes here")
        result = _sample_result(source)

        target = tmp_path / "elsewhere" / "document_ocr.pdf"
        written = ExportManager().export_pdf(result, target)
        assert written == target
        assert target.exists()
        assert target.read_bytes() == source.read_bytes()
        # Source is still intact
        assert source.exists()

    def test_same_path_is_noop(self, tmp_path: Path) -> None:
        """Save As to the exact existing path returns the path without error."""
        from src.application.export_manager import ExportManager

        source = tmp_path / "searchable.pdf"
        source.write_bytes(b"%PDF-1.7\n")
        result = _sample_result(source)

        # mtime before; should survive the no-op
        before = source.stat().st_mtime_ns
        written = ExportManager().export_pdf(result, source)
        after = source.stat().st_mtime_ns
        assert written == source
        assert before == after

    def test_missing_source_raises(self, tmp_path: Path) -> None:
        """A job whose output no longer exists produces an ExportError."""
        from src.application.export_manager import ExportError, ExportManager

        result = _sample_result(tmp_path / "gone.pdf")
        with pytest.raises(ExportError):
            ExportManager().export_pdf(result, tmp_path / "target.pdf")

    def test_creates_parent_dir(self, tmp_path: Path) -> None:
        from src.application.export_manager import ExportManager

        source = tmp_path / "searchable.pdf"
        source.write_bytes(b"%PDF-1.7\n")
        result = _sample_result(source)
        target = tmp_path / "a" / "b" / "c" / "out.pdf"
        ExportManager().export_pdf(result, target)
        assert target.exists()


class TestExportDispatcher:
    def test_pdf_dispatcher_copies(self, tmp_path: Path) -> None:
        """ExportFormat.PDF through the dispatcher now copies the file."""
        source = tmp_path / "searchable.pdf"
        source.write_bytes(b"%PDF-1.7\n")
        result = _sample_result(source)
        target = tmp_path / "copy.pdf"
        written = ExportManager().export(result, target, ExportFormat.PDF)
        assert written == target
        assert target.exists()
        assert target.read_bytes() == source.read_bytes()

    def test_txt_via_dispatcher(self, tmp_path: Path) -> None:
        result = _sample_result(tmp_path / "out.pdf")
        target = tmp_path / "out.txt"
        written = ExportManager().export(result, target, ExportFormat.TXT)
        assert written == target
        assert "Hello world." in target.read_text(encoding="utf-8")

    def test_docx_via_dispatcher(self, tmp_path: Path) -> None:
        pytest.importorskip("docx")
        result = _sample_result(tmp_path / "out.pdf")
        target = tmp_path / "out.docx"
        written = ExportManager().export(result, target, ExportFormat.DOCX)
        assert written == target
        assert target.exists()

    def test_clipboard_without_qapp_raises(self, tmp_path: Path) -> None:
        """Clipboard export requires a QApplication. When no QApp is
        running, it must raise RuntimeError. When a QApp IS running
        (e.g. pytest-qt already created one for GUI tests earlier in
        this session), the export must succeed instead — verify both
        shapes so the test works in any CI matrix."""
        from PySide6.QtWidgets import QApplication

        result = _sample_result(tmp_path / "out.pdf")

        if QApplication.instance() is not None:
            # QApp is alive (pytest-qt session) — clipboard should work.
            ExportManager().export(
                result, tmp_path / "dummy", ExportFormat.CLIPBOARD
            )
        else:
            # No QApp — must raise.
            with pytest.raises(RuntimeError):
                ExportManager().export(
                    result, tmp_path / "dummy", ExportFormat.CLIPBOARD
                )


class TestConcatenateText:
    def test_includes_all_pages(self, tmp_path: Path) -> None:
        result = _sample_result(tmp_path / "out.pdf")
        text = ExportManager._concatenate_text(result)
        assert "Первая страница." in text
        assert "Hello world." in text
        assert "[ERROR: Timeout]" in text
        assert "Page 1" in text and "Page 3" in text
