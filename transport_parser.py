"""Парсер транспортных накладных из PDF в Excel.

Простое десктопное приложение на Python + Tkinter для извлечения данных
российских транспортных накладных из машиночитаемых PDF-файлов (с текстовым
слоем, без OCR) и записи результата в форматированный файл .xlsx.

Зависимости:
    pip install pymupdf openpyxl

Запуск:
    python transport_parser.py
"""
from __future__ import annotations

import os
import re
import sys
import queue
import threading
import traceback
from concurrent.futures import ThreadPoolExecutor, as_completed
from dataclasses import dataclass, field
from functools import lru_cache
from pathlib import Path
from typing import List, Optional, Tuple

import tkinter as tk
from tkinter import ttk, filedialog, messagebox, scrolledtext

try:
    import fitz  # PyMuPDF
except ImportError as exc:  # pragma: no cover
    raise SystemExit(
        "Требуется PyMuPDF. Установите его командой: pip install pymupdf"
    ) from exc

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
except ImportError as exc:  # pragma: no cover
    raise SystemExit(
        "Требуется openpyxl. Установите его командой: pip install openpyxl"
    ) from exc


# ---------------------------------------------------------------------------
# Константы
# ---------------------------------------------------------------------------

APP_TITLE = "Парсер транспортных накладных"
OUTPUT_FILENAME = "extraction.xlsx"
SHEET_NAME = "Extraction"
LOW_TEXT_THRESHOLD = 200  # если меньше — ставим LOW_TEXT
MAX_WORKERS = max(2, (os.cpu_count() or 2))
NOT_FOUND = "отсутствует"
UNREADABLE = "неразборчиво"

COLUMNS: List[Tuple[str, int]] = [
    ("Транспортная накладная", 30),
    ("Дата", 14),
    ("№", 15),
    ("Грузоотправитель", 35),
    ("Груз", 35),
    ("Перевозчик", 35),
    ("Транспортное средство", 22),
    ("Прием груза", 40),
    ("Источник файл", 25),
    ("Примечание", 18),
]


# ---------------------------------------------------------------------------
# Предкомпилированные регулярные выражения (кэш на уровне модуля)
# ---------------------------------------------------------------------------

_FLAGS = re.IGNORECASE | re.UNICODE

RE_WHITESPACE = re.compile(r"\s+", _FLAGS)

# Номер документа (по приоритету)
RE_NUMBER_PATTERNS = tuple(
    re.compile(p, _FLAGS) for p in (
        r"транспортн(?:ая|ой)\s+накладн(?:ая|ой)\s*№\s*([^\n\r]{1,50})",
        r"\bнакладн(?:ая|ой)\b\s*№\s*([^\n\r]{1,50})",
        r"(?:№|\bN\b)\s*([A-Za-zА-Яа-я0-9\-_/]+)",
    )
)

RE_DATE_LABELED = re.compile(r"дата\s*[:№N\-–— ]*\s*(\d{2}\.\d{2}\.\d{4})", _FLAGS)
RE_DATE_ANY = re.compile(r"(\d{2}\.\d{2}\.\d{4})", _FLAGS)

RE_SHIPPER = re.compile(r"грузоотправитель\s*[:\-–—]?\s*([^\n\r]{1,250})", _FLAGS)

RE_CARGO_PATTERNS = tuple(
    re.compile(p, _FLAGS) for p in (
        r"наименовани(?:е|я)\s+груз(?:а|ов)\s*[:\-–—]?\s*([^\n\r]{1,400})",
        r"\bгруз\b\s*[:\-–—]?\s*([^\n\r]{1,400})",
    )
)

RE_CARRIER = re.compile(r"перевозчик\s*[:\-–—]?\s*([^\n\r]{1,250})", _FLAGS)

RE_VEHICLE_PATTERNS = tuple(
    re.compile(p, _FLAGS) for p in (
        r"гос\.?\s*номер\s*[:\-–—]?\s*([^\n\r]{1,40})",
        r"государственн\w*\s+регистрационн\w*\s+номер\s*[:\-–—]?\s*([^\n\r]{1,40})",
        r"рег\.?\s*знак\s*[:\-–—]?\s*([^\n\r]{1,40})",
        r"транспортн\w*\s+средств\w*\s*[:\-–—]?\s*([^\n\r]{1,80})",
    )
)

RE_RECEIVING_HEADER = re.compile(r"при[ёе]м\s+груз\w*", _FLAGS)
RE_RECEIVING_STOP = re.compile(
    r"(сдач\w*\s+груз\w*|выдач\w*\s+груз\w*|доставк\w*\s+груз\w*|отметк\w*)",
    _FLAGS,
)

RE_GARBAGE_CHECK = re.compile(r"[А-Яа-яЁё0-9]", _FLAGS)


# ---------------------------------------------------------------------------
# Извлечение и нормализация текста
# ---------------------------------------------------------------------------

def extract_text_from_pdf(pdf_path: str) -> str:
    """Открывает PDF, извлекает текст со всех страниц. OCR не используется."""
    doc = fitz.open(pdf_path)
    try:
        parts = []
        for page in doc:
            text = page.get_text("text") or ""
            parts.append(text)
        return "\n".join(parts)
    finally:
        doc.close()


def normalize_text(text: str) -> str:
    """Схлопывает все последовательности пробелов/переносов в один пробел."""
    return RE_WHITESPACE.sub(" ", text).strip()


def is_garbage(s: str) -> bool:
    """Нет ни кириллической буквы, ни цифры → считаем мусором."""
    return not RE_GARBAGE_CHECK.search(s)


def _clean_field(value: Optional[str]) -> str:
    """Возвращает очищенное значение или маркер «отсутствует»/«неразборчиво»."""
    if value is None:
        return NOT_FOUND
    value = value.strip().strip(":-–—").strip()
    if not value:
        return NOT_FOUND
    if is_garbage(value):
        return UNREADABLE
    return value


def _first_match(text: str, patterns) -> Optional[str]:
    """Возвращает первую непустую группу из списка regex."""
    for pat in patterns:
        m = pat.search(text)
        if m:
            return m.group(1)
    return None


# ---------------------------------------------------------------------------
# Парсинг полей
# ---------------------------------------------------------------------------

@dataclass
class ParsedRow:
    source: str
    bill_title: str = ""
    date: str = NOT_FOUND
    number: str = NOT_FOUND
    shipper: str = NOT_FOUND
    cargo: str = NOT_FOUND
    carrier: str = NOT_FOUND
    vehicle: str = NOT_FOUND
    receiving: str = NOT_FOUND
    note: str = ""

    def to_row(self) -> List[str]:
        return [
            self.bill_title,
            self.date,
            self.number,
            self.shipper,
            self.cargo,
            self.carrier,
            self.vehicle,
            self.receiving,
            self.source,
            self.note,
        ]


def parse_number(text: str) -> str:
    return _clean_field(_first_match(text, RE_NUMBER_PATTERNS))


def parse_date(text: str) -> str:
    m = RE_DATE_LABELED.search(text)
    if m:
        return m.group(1)
    m = RE_DATE_ANY.search(text)
    if m:
        return m.group(1)
    return NOT_FOUND


def parse_shipper(text: str) -> str:
    m = RE_SHIPPER.search(text)
    return _clean_field(m.group(1) if m else None)


def parse_cargo(text: str) -> str:
    return _clean_field(_first_match(text, RE_CARGO_PATTERNS))


def parse_carrier(text: str) -> str:
    m = RE_CARRIER.search(text)
    return _clean_field(m.group(1) if m else None)


def parse_vehicle(text: str) -> str:
    return _clean_field(_first_match(text, RE_VEHICLE_PATTERNS))


def parse_receiving(text: str) -> str:
    m = RE_RECEIVING_HEADER.search(text)
    if not m:
        return NOT_FOUND
    start = m.end()
    chunk = text[start:start + 1200]
    stop = RE_RECEIVING_STOP.search(chunk)
    if stop:
        chunk = chunk[:stop.start()]
    return _clean_field(chunk)


def parse_document(text: str, source_name: str) -> ParsedRow:
    """Парсит весь документ и возвращает строку результата."""
    normalized = normalize_text(text)

    number = parse_number(normalized)
    row = ParsedRow(
        source=source_name,
        date=parse_date(normalized),
        number=number,
        shipper=parse_shipper(normalized),
        cargo=parse_cargo(normalized),
        carrier=parse_carrier(normalized),
        vehicle=parse_vehicle(normalized),
        receiving=parse_receiving(normalized),
    )

    if number and number != NOT_FOUND and number != UNREADABLE:
        row.bill_title = f"Транспортная накладная № {number}"
    else:
        row.bill_title = "Транспортная накладная"

    if len(normalized) < LOW_TEXT_THRESHOLD:
        row.note = "LOW_TEXT"

    return row


def process_pdf(pdf_path: str) -> Tuple[ParsedRow, Optional[str]]:
    """Обрабатывает один PDF. Возвращает (строку, текст_ошибки_или_None)."""
    name = os.path.basename(pdf_path)
    try:
        text = extract_text_from_pdf(pdf_path)
        row = parse_document(text, name)
        return row, None
    except Exception as exc:  # pragma: no cover - safety net
        err = f"{type(exc).__name__}: {exc}"
        row = ParsedRow(source=name, bill_title="Транспортная накладная")
        row.note = f"ERROR: {err}"
        return row, err


# ---------------------------------------------------------------------------
# Запись Excel
# ---------------------------------------------------------------------------

@lru_cache(maxsize=1)
def _header_style():
    return {
        "font": Font(name="Arial", size=11, bold=True, color="FFFFFFFF"),
        "fill": PatternFill("solid", fgColor="FF4472C4"),
        "alignment": Alignment(
            horizontal="center", vertical="center", wrap_text=True
        ),
        "border": Border(*(Side(style="thin", color="FF000000"),) * 4),
    }


@lru_cache(maxsize=1)
def _data_style():
    return {
        "font": Font(name="Arial", size=10),
        "alignment": Alignment(
            horizontal="left", vertical="top", wrap_text=True
        ),
        "border": Border(*(Side(style="thin", color="FF000000"),) * 4),
    }


def write_excel(rows: List[ParsedRow], output_path: str) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_NAME

    headers = [c[0] for c in COLUMNS]
    ws.append(headers)

    hs = _header_style()
    for col_idx, _ in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = hs["font"]
        cell.fill = hs["fill"]
        cell.alignment = hs["alignment"]
        cell.border = hs["border"]

    ds = _data_style()
    for row_idx, row in enumerate(rows, start=2):
        for col_idx, value in enumerate(row.to_row(), start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.font = ds["font"]
            cell.alignment = ds["alignment"]
            cell.border = ds["border"]

    for col_idx, (_, width) in enumerate(COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    ws.freeze_panes = "A2"
    last_col = get_column_letter(len(headers))
    last_row = max(1, len(rows) + 1)
    ws.auto_filter.ref = f"A1:{last_col}{last_row}"

    wb.save(output_path)


# ---------------------------------------------------------------------------
# GUI
# ---------------------------------------------------------------------------

@dataclass
class WorkerMessage:
    kind: str  # "log", "progress", "done", "error"
    payload: object = None


class ParserApp:
    def __init__(self, root: tk.Tk) -> None:
        self.root = root
        self.root.title(APP_TITLE)
        self.root.geometry("720x540")
        self.root.minsize(640, 480)

        self.input_var = tk.StringVar()
        self.output_var = tk.StringVar()
        self.queue: "queue.Queue[WorkerMessage]" = queue.Queue()
        self.worker: Optional[threading.Thread] = None

        self._build_ui()
        self.root.after(80, self._drain_queue)

    # ---- UI --------------------------------------------------------------

    def _build_ui(self) -> None:
        pad = {"padx": 10, "pady": 6}

        frm_in = ttk.Frame(self.root)
        frm_in.pack(fill="x", **pad)
        ttk.Label(frm_in, text="Папка с PDF:", width=14).pack(side="left")
        ttk.Entry(frm_in, textvariable=self.input_var).pack(
            side="left", fill="x", expand=True, padx=(0, 6)
        )
        ttk.Button(frm_in, text="Обзор…", command=self._pick_input).pack(side="left")

        frm_out = ttk.Frame(self.root)
        frm_out.pack(fill="x", **pad)
        ttk.Label(frm_out, text="Сохранить в:", width=14).pack(side="left")
        ttk.Entry(frm_out, textvariable=self.output_var).pack(
            side="left", fill="x", expand=True, padx=(0, 6)
        )
        ttk.Button(frm_out, text="Обзор…", command=self._pick_output).pack(side="left")

        frm_btn = ttk.Frame(self.root)
        frm_btn.pack(**pad)
        self.run_btn = ttk.Button(
            frm_btn, text="▶  Извлечь данные", command=self._start
        )
        self.run_btn.pack()

        ttk.Label(self.root, text="Лог:").pack(anchor="w", padx=10)
        self.log = scrolledtext.ScrolledText(
            self.root, height=14, state="disabled", wrap="word"
        )
        self.log.pack(fill="both", expand=True, padx=10, pady=(0, 6))

        frm_pb = ttk.Frame(self.root)
        frm_pb.pack(fill="x", padx=10, pady=(0, 10))
        ttk.Label(frm_pb, text="Прогресс:").pack(side="left", padx=(0, 6))
        self.progress = ttk.Progressbar(frm_pb, mode="determinate")
        self.progress.pack(side="left", fill="x", expand=True)
        self.progress_label = ttk.Label(frm_pb, text="0 / 0", width=10)
        self.progress_label.pack(side="left", padx=(6, 0))

    # ---- actions ---------------------------------------------------------

    def _pick_input(self) -> None:
        path = filedialog.askdirectory(title="Папка с PDF")
        if path:
            self.input_var.set(path)
            if not self.output_var.get():
                self.output_var.set(path)

    def _pick_output(self) -> None:
        path = filedialog.askdirectory(title="Папка для сохранения")
        if path:
            self.output_var.set(path)

    def _append_log(self, text: str) -> None:
        self.log.configure(state="normal")
        self.log.insert("end", text + "\n")
        self.log.see("end")
        self.log.configure(state="disabled")

    def _set_running(self, running: bool) -> None:
        self.run_btn.configure(state="disabled" if running else "normal")

    def _start(self) -> None:
        in_dir = self.input_var.get().strip()
        out_dir = self.output_var.get().strip() or in_dir

        if not in_dir or not os.path.isdir(in_dir):
            messagebox.showwarning(APP_TITLE, "Выберите папку с PDF.")
            return
        if not out_dir:
            out_dir = in_dir
        if not os.path.isdir(out_dir):
            messagebox.showwarning(APP_TITLE, "Папка для сохранения не существует.")
            return

        pdfs = sorted(
            str(p) for p in Path(in_dir).iterdir()
            if p.is_file() and p.suffix.lower() == ".pdf"
        )
        if not pdfs:
            messagebox.showwarning(APP_TITLE, "В выбранной папке нет PDF-файлов.")
            return

        self.log.configure(state="normal")
        self.log.delete("1.0", "end")
        self.log.configure(state="disabled")
        self.progress["value"] = 0
        self.progress["maximum"] = len(pdfs)
        self.progress_label.configure(text=f"0 / {len(pdfs)}")

        self._set_running(True)
        self._append_log(f"Найдено PDF-файлов: {len(pdfs)}")

        output_path = os.path.join(out_dir, OUTPUT_FILENAME)
        self.worker = threading.Thread(
            target=self._run_worker, args=(pdfs, output_path), daemon=True
        )
        self.worker.start()

    # ---- worker ----------------------------------------------------------

    def _run_worker(self, pdfs: List[str], output_path: str) -> None:
        rows: List[ParsedRow] = []
        ok = 0
        errors = 0
        # Пределяем порядок по исходному списку, несмотря на параллелизм.
        index_by_path = {p: i for i, p in enumerate(pdfs)}
        results: List[Optional[Tuple[ParsedRow, Optional[str]]]] = [None] * len(pdfs)

        try:
            with ThreadPoolExecutor(max_workers=MAX_WORKERS) as pool:
                futures = {pool.submit(process_pdf, p): p for p in pdfs}
                processed = 0
                for fut in as_completed(futures):
                    path = futures[fut]
                    name = os.path.basename(path)
                    try:
                        row, err = fut.result()
                    except Exception as exc:  # pragma: no cover
                        err = f"{type(exc).__name__}: {exc}"
                        row = ParsedRow(
                            source=name,
                            bill_title="Транспортная накладная",
                            note=f"ERROR: {err}",
                        )
                    results[index_by_path[path]] = (row, err)
                    processed += 1
                    if err is None:
                        ok += 1
                        self.queue.put(WorkerMessage("log", f"Обработано: {name} — OK"))
                    else:
                        errors += 1
                        self.queue.put(
                            WorkerMessage("log", f"Ошибка: {name} — {err}")
                        )
                    self.queue.put(WorkerMessage("progress", processed))

            rows = [r[0] for r in results if r is not None]
            write_excel(rows, output_path)
            self.queue.put(WorkerMessage(
                "log",
                f"──────────────────────────\n"
                f"Итого: {len(pdfs)} файлов, {ok} OK, {errors} ошибка(и)\n"
                f"Сохранено: {output_path}",
            ))
            self.queue.put(WorkerMessage("done"))
        except Exception:
            tb = traceback.format_exc()
            self.queue.put(WorkerMessage("error", tb))

    # ---- main-thread queue drain ----------------------------------------

    def _drain_queue(self) -> None:
        try:
            while True:
                msg = self.queue.get_nowait()
                if msg.kind == "log":
                    self._append_log(str(msg.payload))
                elif msg.kind == "progress":
                    done = int(msg.payload or 0)
                    self.progress["value"] = done
                    total = int(self.progress["maximum"])
                    self.progress_label.configure(text=f"{done} / {total}")
                elif msg.kind == "done":
                    self._set_running(False)
                elif msg.kind == "error":
                    self._append_log("Критическая ошибка:\n" + str(msg.payload))
                    self._set_running(False)
        except queue.Empty:
            pass
        finally:
            self.root.after(80, self._drain_queue)


# ---------------------------------------------------------------------------
# entry point
# ---------------------------------------------------------------------------

def main() -> int:
    root = tk.Tk()
    try:
        style = ttk.Style(root)
        if "vista" in style.theme_names():
            style.theme_use("vista")
        elif "clam" in style.theme_names():
            style.theme_use("clam")
    except tk.TclError:
        pass
    ParserApp(root)
    root.mainloop()
    return 0


if __name__ == "__main__":
    sys.exit(main())
